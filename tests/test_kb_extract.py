import io

import pytest

from app.kb.extract import (
    DOCX_MIME,
    PPTX_MIME,
    XLSX_MIME,
    ExtractionError,
    extract_text,
    is_supported,
)


def _make_pdf(text: str) -> bytes:
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R "
        b"/Resources << /Font << /F1 5 0 R >> >> >>",
    ]
    stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
    objects.append(b"<< /Length %d >>\nstream\n" % len(stream) + stream + b"\nendstream")
    objects.append(b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>")
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % number + body + b"\nendobj\n"
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    for offset in offsets:
        out += b"%010d 00000 n \n" % offset
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (len(objects) + 1, xref)
    return out


def test_supported_types() -> None:
    assert is_supported("application/pdf")
    assert is_supported("application/vnd.google-apps.document")
    assert is_supported("text/markdown")
    assert not is_supported("image/png")
    assert not is_supported("application/vnd.google-apps.shortcut")


def test_markdown_and_plain_text() -> None:
    assert extract_text("# Titre\n\nCorps  \n".encode(), "text/markdown") == "# Titre\n\nCorps"


def test_docx_keeps_headings_paragraphs_and_tables() -> None:
    from docx import Document

    document = Document()
    document.add_heading("Politique mobile", level=1)
    document.add_paragraph("Les appareils doivent être chiffrés.")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Appareil"
    table.cell(0, 1).text = "Règle"
    table.cell(1, 0).text = "Mobile"
    table.cell(1, 1).text = "MDM obligatoire"
    buffer = io.BytesIO()
    document.save(buffer)

    text = extract_text(buffer.getvalue(), DOCX_MIME)
    assert "# Politique mobile" in text
    assert "Les appareils doivent être chiffrés." in text
    assert "Mobile | MDM obligatoire" in text


def test_pptx_extracts_slide_text_and_notes() -> None:
    from pptx import Presentation

    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[1])
    slide.shapes.title.text = "Serveur MCP"
    slide.placeholders[1].text = "Couche de minimisation des données"
    slide.notes_slide.notes_text_frame.text = "Insister sur la portabilité"
    buffer = io.BytesIO()
    presentation.save(buffer)

    text = extract_text(buffer.getvalue(), PPTX_MIME)
    assert "## Diapositive 1" in text
    assert "Serveur MCP" in text
    assert "Couche de minimisation des données" in text
    assert "Notes : Insister sur la portabilité" in text


def test_xlsx_reads_sheets_as_rows() -> None:
    from openpyxl import Workbook

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Budget"
    sheet.append(["Poste", "Montant"])
    sheet.append(["Licences", 1200])
    buffer = io.BytesIO()
    workbook.save(buffer)

    text = extract_text(buffer.getvalue(), XLSX_MIME)
    assert "## Feuille Budget" in text
    assert "Licences | 1200" in text


def test_pdf_with_text() -> None:
    phrase = "Le budget de la campagne est valide par la direction de la communication. " * 5
    text = extract_text(_make_pdf(phrase), "application/pdf")
    assert "[page 1]" in text
    assert "budget de la campagne" in text


def test_pdf_without_text_is_reported_as_scanned() -> None:
    from pypdf import PdfWriter

    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    with pytest.raises(ExtractionError, match="scanné"):
        extract_text(buffer.getvalue(), "application/pdf")


def test_unsupported_and_empty_and_corrupt() -> None:
    with pytest.raises(ExtractionError, match="non pris en charge"):
        extract_text(b"x", "image/png")
    with pytest.raises(ExtractionError, match="aucun texte"):
        extract_text(b"   \n  ", "text/plain")
    with pytest.raises(ExtractionError, match="lecture impossible"):
        extract_text(b"pas un vrai docx", DOCX_MIME)
